"""Page/slide/sheet/frame enumeration. Optional format libraries are loaded lazily."""
from __future__ import annotations

import csv
import importlib.util
import shutil
import subprocess
import tempfile
from pathlib import Path
from .schemas import Document, FileType, Page


class LoaderFailure(Exception):
    def __init__(self, code: str, message: str): super().__init__(message); self.code = code


def load_pages(document: Document) -> list[Page]:
    p = Path(document.source); kind = document.file_type
    if kind == FileType.PDF: return _load_pdf(p)
    if kind in (FileType.PNG, FileType.JPG, FileType.TIFF): return _load_image(p)
    if kind == FileType.TXT:
        separators=0
        with p.open(encoding="utf-8-sig") as stream:
            for chunk in iter(lambda:stream.read(1024*1024),""): separators+=chunk.count("\f")
        return [Page(str(p),i+1,"page",metadata={"text_source":str(p)}) for i in range(separators+1)]
    if kind == FileType.CSV:
        row_count=column_count=0
        with p.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.reader(f): row_count+=1; column_count=max(column_count,len(row))
        return [Page(str(p), 1, "sheet", metadata={"row_count":row_count,"column_count":column_count})]
    if kind == FileType.DOCX: return _load_docx(p)
    if kind == FileType.XLSX: return _load_xlsx(p)
    if kind == FileType.PPTX: return _load_pptx(p)
    if kind in (FileType.DOC, FileType.XLS, FileType.PPT): return _load_legacy(p, kind)
    raise LoaderFailure("UNSUPPORTED_FORMAT", f"No loader for {kind.value}")


def _require(module: str, extra: str):
    if importlib.util.find_spec(module) is None:
        raise LoaderFailure("MISSING_DEPENDENCY", f"Install parseanything[{extra}] to load this format")


def _load_legacy(p, kind):
    """Use LibreOffice headless conversion when available; never modify the source."""
    executable=shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise LoaderFailure("LEGACY_ADAPTER_REQUIRED", f"Install LibreOffice or configure a loader for legacy {kind.value}")
    target={FileType.DOC:"docx",FileType.XLS:"xlsx",FileType.PPT:"pptx"}[kind]
    try:
        with tempfile.TemporaryDirectory(prefix="parseanything-") as temp:
            proc=subprocess.run([executable,"--headless","--convert-to",target,"--outdir",temp,str(p)],
                capture_output=True,text=True,timeout=60,check=False)
            converted=Path(temp)/(p.stem+"."+target)
            if proc.returncode or not converted.exists():
                raise LoaderFailure("LEGACY_CONVERSION_FAILED",(proc.stderr or proc.stdout or "LibreOffice conversion failed").strip())
            if target=="docx": return _load_docx(converted)
            if target=="xlsx": return _load_xlsx(converted)
            return _load_pptx(converted)
    except subprocess.TimeoutExpired as e:
        raise LoaderFailure("TIMEOUT","Legacy Office conversion exceeded 60 seconds") from e


def _load_pdf(p):
    if importlib.util.find_spec("fitz") is None and importlib.util.find_spec("pdfplumber") is None:
        raise LoaderFailure("MISSING_DEPENDENCY", "Install PyMuPDF or pdfplumber to load PDF files")
    try:
        pages = []
        if importlib.util.find_spec("fitz") is not None:
            import fitz
            doc = fitz.open(p)
            if doc.is_encrypted and not doc.authenticate(""): raise LoaderFailure("ENCRYPTED_PDF", "Password-protected PDF")
            for i, page in enumerate(doc):
                raw = page.get_text("dict")
                text_blocks=[]
                for b in raw.get("blocks",[]):
                    if b.get("type") != 0: continue
                    spans=[s for line in b.get("lines",[]) for s in line.get("spans",[])]
                    text_blocks.append({"bbox":b.get("bbox"),"type":b.get("type"),
                        "font_size_min":min((float(s.get("size",99)) for s in spans),default=99),
                        "font_names":sorted({str(s.get("font","")) for s in spans})})
                has_text=bool(text_blocks)
                images = page.get_images(full=True)
                mode = "digital" if has_text else "scanned" if images else "empty"
                pages.append(Page(str(p), i+1, width=float(page.rect.width), height=float(page.rect.height),
                    metadata={"pdf_type": mode, "image_count": len(images), "pdf_engine": "pymupdf"},
                    payload={"page": page, "text_blocks": text_blocks}))
        else:
            import pdfplumber
            pdf = pdfplumber.open(p)
            for i, page in enumerate(pdf.pages):
                has_text = bool((page.extract_text() or "").strip())
                images = page.images or []
                mode = "digital" if has_text else "scanned" if images else "empty"
                pages.append(Page(str(p), i+1, width=float(page.width), height=float(page.height),
                    metadata={"pdf_type": mode, "image_count": len(images), "pdf_engine": "pdfplumber"},
                    payload={"page": page,"has_text":has_text}))
        if not pages: raise LoaderFailure("INVALID_PDF", "PDF has no pages")
        return pages
    except LoaderFailure: raise
    except Exception as e: raise LoaderFailure("CORRUPT_PDF", str(e)) from e


def _load_image(p):
    _require("PIL", "images")
    from PIL import Image
    try:
        with Image.open(p) as im:
            result = [Page(str(p), i+1, "frame", *im.size, metadata={"format": im.format, "mode": im.mode,
                           "image_path": str(p), "frame_index": i})
                      for i in range(getattr(im, "n_frames", 1))]
        return result
    except Exception as e: raise LoaderFailure("CORRUPT_IMAGE", str(e)) from e


def _load_docx(p):
    _require("docx", "office")
    from docx import Document as Docx
    try:
        d = Docx(p)
        return [Page(str(p), 1, "flow", metadata={"pagination": "requires_rendering"},
                     payload=d)]
    except Exception as e: raise LoaderFailure("CORRUPT_DOCX", str(e)) from e


def _load_xlsx(p):
    _require("openpyxl", "office")
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=False, data_only=False)
        pages = [Page(str(p), i+1, "sheet", metadata={"name": s.title, "max_row":s.max_row,
                    "max_column":s.max_column,"merged_ranges": [str(r) for r in s.merged_cells.ranges]}, payload=s)
                 for i, s in enumerate(wb.worksheets)]
        return pages
    except Exception as e: raise LoaderFailure("CORRUPT_XLSX", str(e)) from e


def _load_pptx(p):
    _require("pptx", "office")
    from pptx import Presentation
    try:
        d = Presentation(p)
        return [Page(str(p), i+1, "slide", d.slide_width/914400, d.slide_height/914400,
                     payload=slide) for i, slide in enumerate(d.slides)]
    except Exception as e: raise LoaderFailure("CORRUPT_PPTX", str(e)) from e
